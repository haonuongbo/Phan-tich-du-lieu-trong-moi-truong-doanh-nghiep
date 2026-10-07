"""Tạo báo cáo và bản xem dashboard offline từ cùng một bộ KPI."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from pipeline import ROOT, LABEL


def number(value, digits=2):
    if pd.isna(value):
        return "Chưa đủ dữ liệu"
    return f"{value:,.{digits}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def percent(value):
    return number(value * 100) + "%" if pd.notna(value) else "Chưa đủ dữ liệu"


def report_content(tables, quality):
    annual = tables["KPI_Nam"]
    latest = annual.iloc[-1]
    stores = tables["KPI_CuaHang_Nam"].query("nam == @latest.nam").sort_values("DT_tong", ascending=False)
    monthly = tables["KPI_Thang"].query("nam == @latest.nam")
    peak = monthly.loc[monthly.DT_tong.idxmax()]
    bottom = stores.sort_values("bien_LN_van_hanh").iloc[0]
    loss = stores[stores.LN_van_hanh_mo_hinh.lt(0)]
    cagr = tables["KPI_TongQuan"].iloc[0].CAGR_doanh_thu
    sections = [
        ("1. Mục tiêu và phạm vi", [
            "Thiết kế hệ thống phân tích doanh thu KFC tại TP.HCM theo tài liệu THIẾT KẾ HỆ THỐNG.docx: quản lý nguồn → làm sạch bằng Pandas → tính KPI → cung cấp dữ liệu Power BI → báo cáo và đề xuất.",
            f"Nguồn cửa hàng: 02_KFC_TPHCM_cua_hang_mo_phong.csv. Nguồn vận hành: {quality['nguon_van_hanh']}. Phạm vi {quality['tu_thang']} đến {quality['den_thang']}; {quality['so_dong_sach']} bản ghi cửa hàng–tháng, {quality['so_cua_hang']} cửa hàng hư cấu.",
            "Nguồn vận hành DULIEUTHO.csv là CSV UTF-8 đúng định dạng, chuyển từ dữ liệu Excel được cung cấp. Hai nguồn thô có lỗi mô phỏng để thực hành làm sạch; không lấy KPI có sẵn làm nguồn."]),
        ("2. Kiến trúc và xử lý dữ liệu", [
            "Pipeline đọc hai nguồn thô trong DataGoc. DataSach/DULIEUSACH.csv là một tệp sạch duy nhất, gộp vận hành và thuộc tính cửa hàng. DataPhantich chứa mô hình sao và bảng KPI dùng cho dashboard.",
            "Mỗi dòng Fact_Operation là một cửa hàng trong một tháng; khóa ghép store_id + thang. Fact_Channel có thêm khóa kenh. Chiều thời gian liên tục theo ngày, quan hệ một–nhiều, lọc một chiều từ dimension sang fact.",
            f"Nguồn có {quality['so_dong_nguon_cua_hang']} dòng cửa hàng và {quality['so_dong_nguon']} dòng vận hành. Loại {quality['cua_hang_trung_hoan_toan_da_bo']} + {quality['van_hanh_trung_hoan_toan_da_bo']} dòng trùng và {quality['so_dong_loi_cua_hang']} + {quality['so_dong_loi_van_hanh']} dòng lỗi không thể phục hồi. Xem nhat_ky_lam_sach.csv và dong_bi_loai.csv.",
            "CSV ngày nguồn dùng tháng/ngày/năm, chuẩn hóa thành YYYY-MM-DD. Tiền là VND; tỷ lệ lưu 0–1. Tính lại doanh thu = số đơn × AOV từng kênh; đối chiếu với cột có sẵn, dung sai 1 VND và 1e-8 cho tỷ lệ do làm tròn nguồn.",
            f"Chất lượng sau làm sạch: {quality['trang_thai']}. Đối chiếu {quality['so_dong_doi_chieu_cong_thuc']} dòng sạch; cột dẫn xuất sai hoặc thiếu được tính lại từ số đơn, AOV và chi phí. Có {quality['so_o_da_sua']} lượt sửa/chuẩn hóa ô có nhật ký; không điền số 0 thay cho dữ liệu thiếu."]),
        ("3. Định nghĩa KPI", [
            "Doanh thu = tổng doanh thu ba kênh. Tăng trưởng kỳ = doanh thu kỳ này / kỳ trước − 1, chỉ tính khi hai kỳ đủ tháng. YoY so cùng kỳ năm trước. CAGR = (tổng doanh thu năm cuối / năm đầu)^(1 / khoảng cách năm) − 1; chỉ dùng năm đủ 12 tháng.",
            "Cửa hàng bình quân = số bản ghi cửa hàng–tháng / số tháng trong kỳ. Doanh thu/cửa hàng = tổng doanh thu / cửa hàng bình quân. Đồng thời xuất số cửa hàng đầu kỳ, cuối kỳ và trung bình hai mốc để đối chiếu; không cộng số cửa hàng qua các tháng.",
            "Lợi nhuận gộp = doanh thu − giá vốn. Lợi nhuận vận hành mô hình = lợi nhuận gộp − phí nền tảng − phí nhượng quyền − lương − thuê − chi phí khác − khấu hao. Biên vận hành = tổng lợi nhuận vận hành / tổng doanh thu.",
            "AOV = tổng doanh thu / tổng đơn. Tỷ trọng online = (doanh thu nền tảng + online riêng) / tổng doanh thu. Không lấy trung bình cộng AOV hoặc biên lợi nhuận của các dòng.",
            "SSSG chỉ so cùng tập cửa hàng có đủ 12 tháng ở cả hai năm. Thiếu kỳ trước hoặc mẫu số bằng 0: để trống, không suy diễn bằng 0%. Không có lợi nhuận sau thuế, tài sản/vốn, khách hàng định danh, dữ liệu ngành nên chưa tính biên ròng, ROA/ROE, tỷ lệ quay lại hoặc thị phần."]),
        (f"4. Kết quả mô phỏng năm {int(latest.nam)}", [
            f"Doanh thu {number(latest.DT_tong / 1e9)} tỷ VND; {number(latest.tong_don, 0)} đơn; cửa hàng bình quân {number(latest.CH_binh_quan, 0)}; doanh thu/cửa hàng {number(latest.DT_CH / 1e9)} tỷ VND/năm.",
            f"Lợi nhuận vận hành mô hình {number(latest.LN_van_hanh_mo_hinh / 1e9)} tỷ VND; biên vận hành {percent(latest.bien_LN_van_hanh)}; AOV {number(latest.AOV_tong, 0)} VND/đơn; tỷ trọng doanh thu online {percent(latest.ty_trong_DT_online)}.",
            f"YoY doanh thu: {percent(latest.YoY_doanh_thu)}. CAGR toàn khoảng đủ năm: {percent(cagr)}. Tăng trưởng cùng cửa hàng: {percent(latest.SSSG)}.",
            f"Tháng có doanh thu cao nhất năm: {peak.ky:%m/%Y}, đạt {number(peak.DT_tong / 1e9)} tỷ VND. Cửa hàng đóng góp doanh thu cao nhất: {stores.iloc[0].store_id}, đạt {number(stores.iloc[0].DT_tong / 1e9)} tỷ VND.",
            f"Có {len(loss)} cửa hàng lỗ vận hành cả năm. Biên vận hành thấp nhất: {bottom.store_id} ({percent(bottom.bien_LN_van_hanh)}). Đây là tín hiệu mô tả, chưa chứng minh nguyên nhân."]),
        ("5. Định hướng quyết định trong bài mô phỏng", [
            f"Ưu tiên rà soát cơ cấu giá vốn, thuê mặt bằng và nhân sự tại {bottom.store_id}; so doanh thu và từng khoản chi phí theo tháng trước khi chọn biện pháp. Theo dõi số cửa hàng lỗ và biên vận hành ở kỳ tiếp theo.",
            "Tách nền tảng giao hàng và online riêng khi xem đóng góp doanh thu; cần phân bổ thêm chi phí theo kênh trước khi kết luận kênh nào sinh lợi cao hơn. Chưa có cơ sở nhân quả để khẳng định tăng online làm tăng lợi nhuận.",
            "Chưa có kế hoạch KPI được phê duyệt nên trạng thái chỉ mô tả biến động, không tự gán đạt/không đạt mục tiêu. Chuỗi 2023–2025 cho phép so YoY/CAGR và quan sát mùa vụ; chưa có dữ liệu ngành cùng phạm vi để so vị thế thị trường."]),
        ("6. Dashboard và vận hành", [
            "Gói PowerBI gồm hàm Power Query, mã DAX, theme và hướng dẫn bốn trang: tổng quan, xu hướng, hiệu quả cửa hàng, chi phí/lợi nhuận. Người dùng nhập vào Power BI Desktop và lưu PBIX; gói này chưa chứa tệp PBIX đã dựng hoặc đã kiểm tra bằng Power BI Desktop.",
            "BaoCao/dashboard.html là bản xem tương tác độc lập có lọc năm/quý/cửa hàng và xuất CSV theo bộ lọc; mở trực tiếp trong trình duyệt, không cần mạng. Đây là bản xem bổ trợ cho việc kiểm tra và thuyết trình.",
            "Chạy lại Script/Main.py khi cập nhật nguồn. Đọc chat_luong_du_lieu.json trước khi Refresh Power BI. Nếu trạng thái LOI, đầu ra cũ không đại diện cho lần cập nhật thất bại. Chạy bộ kiểm thử trong Tests trước khi thay đổi công thức."]),
        ("7. Giới hạn kết luận", [
            LABEL,
            "Chưa xác minh cơ sở doanh thu gộp/thuần và thuế VAT của mô hình. Không gọi lợi nhuận vận hành là lợi nhuận ròng. Không suy rộng quy mô 20 cửa hàng hư cấu thành quy mô thực tế của KFC TP.HCM.",
            "Bộ nguồn không có ngày đóng cửa hay quận/huyện; giả định cửa hàng hoạt động liên tục từ tháng mở. Khi có đóng cửa hoặc hoạt động giữa tháng cần mở rộng mô hình và quy tắc tính cửa hàng bình quân.",
            "Doanh thu tăng, nếu có, chưa đủ kết luận hiệu quả kinh doanh. Cần đối chiếu doanh thu/cửa hàng, biên vận hành, cùng tập cửa hàng và thị trường có phạm vi tương đương."])
    ]
    return sections


def generate_dashboard(tables, quality, project=ROOT):
    project = Path(project)
    payload = {"fact": json.loads(tables['Fact_Operation'].to_json(orient='records',date_format='iso')),
               "menu": json.loads(tables['Fact_Menu'].to_json(orient='records',date_format='iso')),
               "stores": json.loads(tables['Dim_Store'].to_json(orient='records',date_format='iso')),
               "quality": quality}
    template = (project / 'Script/dashboard.html').read_text(encoding='utf-8')
    (project / 'BaoCao/dashboard.html').write_text(template.replace('__PAYLOAD__',json.dumps(payload,ensure_ascii=False).replace('</','<\\/')),encoding='utf-8')


def generate(tables, quality, project=ROOT):
    project = Path(project)
    report_dir = project / "BaoCao"
    sections = report_content(tables, quality)
    text = "# HỆ THỐNG PHÂN TÍCH DOANH THU KFC TP.HCM\n\n" + LABEL + "\n\n"
    for title, paragraphs in sections:
        text += "## " + title + "\n\n" + "\n\n".join(paragraphs) + "\n\n"
    (report_dir / "BaoCao_PhanTich.md").write_text(text, encoding="utf-8")
    generate_dashboard(tables, quality, project)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = DocPt(11)
    doc.add_heading("HỆ THỐNG PHÂN TÍCH DOANH THU\nKFC TP.HCM", 0)
    doc.add_paragraph(LABEL, style="Subtitle")
    for title, paragraphs in sections:
        doc.add_heading(title, 1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    doc.add_heading("Phụ lục: đối chiếu KPI năm", 1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, ["Năm", "DT (tỷ VND)", "LN vận hành (tỷ)", "Biên LN", "YoY DT"]):
        cell.text = label
    for r in tables["KPI_Nam"].itertuples():
        for cell, value in zip(table.add_row().cells, [str(r.nam), number(r.DT_tong / 1e9),
            number(r.LN_van_hanh_mo_hinh / 1e9), percent(r.bien_LN_van_hanh), percent(r.YoY_doanh_thu)]):
            cell.text = value
    for section in doc.sections:
        section.header.paragraphs[0].text = "KFC TP.HCM | BÀI THỰC HÀNH DỮ LIỆU MÔ PHỎNG"
        section.footer.paragraphs[0].text = "Nguồn 02 + DULIEUTHO • Dữ liệu mô phỏng"
        section.top_margin = DocInches(0.7)
        section.bottom_margin = DocInches(0.7)
    try:
        doc.save(report_dir / "BaoCao_PhanTich.docx")
    except PermissionError:
        print("Chưa cập nhật báo cáo Word vì tệp đang mở/khóa. Đóng tệp và chạy lại Main.py.")

    slides = Presentation()
    slides.slide_width, slides.slide_height = Inches(13.333), Inches(7.5)
    for title, paragraphs in sections[:6]:
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FAF8F5")
        title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.45), Inches(12), Inches(0.9))
        title_box.text = title
        title_box.text_frame.paragraphs[0].font.size = Pt(30)
        title_box.text_frame.paragraphs[0].font.bold = True
        title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("AC1232")
        body = slide.shapes.add_textbox(Inches(0.75), Inches(1.65), Inches(11.8), Inches(4.9)).text_frame
        body.word_wrap = True
        for index, paragraph in enumerate(paragraphs[:4]):
            p = body.paragraphs[0] if index == 0 else body.add_paragraph()
            p.text = paragraph
            p.font.size = Pt(19 if len(paragraphs) <= 3 else 17)
            p.space_after = Pt(17)
        footer = slide.shapes.add_textbox(Inches(0.75), Inches(7), Inches(12), Inches(0.3))
        footer.text = LABEL
        footer.text_frame.paragraphs[0].font.size = Pt(11)
    chart_slide = slides.slides.add_slide(slides.slide_layouts[6])
    title = chart_slide.shapes.add_textbox(Inches(0.65), Inches(0.4), Inches(12), Inches(0.8))
    title.text = "Doanh thu theo tháng • Dữ liệu mô phỏng"
    title.text_frame.paragraphs[0].font.size = Pt(30)
    latest_year = tables["KPI_Nam"].nam.max()
    monthly = tables["KPI_Thang"].query("nam == @latest_year")
    chart_data = CategoryChartData()
    chart_data.categories = [f"{d:%m/%Y}" for d in monthly.ky]
    chart_data.add_series("Doanh thu (tỷ VND)", list(monthly.DT_tong / 1e9))
    chart = chart_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2), chart_data).chart
    chart.has_legend = False
    chart.chart_style = 10
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Tỷ VND"
    footer = chart_slide.shapes.add_textbox(Inches(0.7), Inches(7), Inches(12), Inches(0.3))
    footer.text = LABEL
    footer.text_frame.paragraphs[0].font.size = Pt(11)
    try:
        slides.save(report_dir / "ThuyetTrinh.pptx")
    except PermissionError:
        print("Chưa cập nhật slide vì tệp đang mở/khóa. Đóng tệp và chạy lại Main.py.")
